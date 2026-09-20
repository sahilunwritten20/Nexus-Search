"""Files connector for Nexus Search.

Supports: TXT, Markdown, RST, CSV, JSON, HTML, PDF, DOCX, XLSX, PPTX.
Unreadable files are logged (never silently dropped) and skipped.
"""
import csv
import json
import logging
from pathlib import Path
from typing import Callable, Iterator, Optional

from ..types import IngestDoc

logger = logging.getLogger("nexus_search.ingestion.files")

DEFAULT_EXTENSIONS = {
    ".txt", ".md", ".rst", ".csv", ".json", ".html", ".htm",
    ".pdf", ".docx", ".xlsx", ".pptx",
}
IGNORED_DIRS = {".git", ".venv", "venv", "node_modules", "__pycache__", "dist", "build"}


def read_text_file(path: Path) -> str:
    return path.read_text(encoding="utf-8", errors="ignore")


def read_csv_file(path: Path) -> str:
    with path.open(newline="", encoding="utf-8", errors="ignore") as f:
        rows = [" | ".join(f"{k}: {v}" for k, v in row.items() if k and v) for row in csv.DictReader(f)]
    return "\n".join(r for r in rows if r)


def read_json_file(path: Path) -> str:
    out: list[str] = []

    def walk(x):
        if isinstance(x, dict):
            for k, v in x.items():
                out.append(str(k))
                walk(v)
        elif isinstance(x, list):
            for v in x:
                walk(v)
        elif x is not None:
            out.append(str(x))

    walk(json.loads(path.read_text(encoding="utf-8", errors="ignore")))
    return "\n".join(out)


def read_html_file(path: Path) -> str:
    from .web import parse_html

    return parse_html(path.read_text(encoding="utf-8", errors="ignore"), path.as_uri()).content


def read_pdf_file(path: Path) -> str:
    from pypdf import PdfReader

    pages = (page.extract_text() for page in PdfReader(str(path)).pages)
    return "\n".join(t for t in pages if t)


def read_docx_file(path: Path) -> str:
    from docx import Document

    document = Document(str(path))
    text = [p.text for p in document.paragraphs if p.text]
    for table in document.tables:  # tables were previously dropped
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                text.append(" | ".join(cells))
    return "\n".join(text)


def read_xlsx_file(path: Path) -> str:
    from openpyxl import load_workbook

    workbook = load_workbook(filename=str(path), read_only=True, data_only=True)
    try:
        text = []
        for sheet in workbook.worksheets:
            text.append(f"Sheet: {sheet.title}")
            for row in sheet.iter_rows(values_only=True):
                values = [str(v) for v in row if v is not None]
                if values:
                    text.append(" | ".join(values))
        return "\n".join(text)
    finally:
        workbook.close()


def read_pptx_file(path: Path) -> str:
    from pptx import Presentation

    text = []
    for n, slide in enumerate(Presentation(str(path)).slides, start=1):
        text.append(f"Slide: {n}")
        for shape in slide.shapes:
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    text.append(" | ".join(c.text for c in row.cells if c.text))
            elif getattr(shape, "text", ""):
                text.append(shape.text)
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = slide.notes_slide.notes_text_frame.text
            if notes:
                text.append(notes)
    return "\n".join(text)


_READERS: dict[str, Callable[[Path], str]] = {
    ".txt": read_text_file, ".md": read_text_file, ".rst": read_text_file,
    ".csv": read_csv_file, ".json": read_json_file,
    ".html": read_html_file, ".htm": read_html_file,
    ".pdf": read_pdf_file, ".docx": read_docx_file,
    ".xlsx": read_xlsx_file, ".pptx": read_pptx_file,
}


def read_file(path: Path) -> str:
    """Read a file according to its extension; '' (and a logged warning) on failure."""
    reader = _READERS.get(path.suffix.lower())
    if reader is None:
        return ""
    try:
        return reader(path)
    except Exception as exc:  # corrupt file, missing optional dependency, ...
        logger.warning("Could not read %s: %s: %s", path, type(exc).__name__, exc)
        return ""


def iter_files(root: str, extensions: Optional[set[str]] = None) -> Iterator[IngestDoc]:
    extensions = {e.lower() for e in (extensions or DEFAULT_EXTENSIONS)}
    root_path = Path(root)

    for path in sorted(root_path.rglob("*")):
        rel = path.relative_to(root_path)
        if not path.is_file() or path.suffix.lower() not in extensions:
            continue
        if any(part in IGNORED_DIRS for part in rel.parts):
            continue
        content = read_file(path)
        if not content.strip():
            logger.warning("No text extracted from %s (empty, scanned, or unreadable)", rel)
            continue
        yield IngestDoc(
            doc_id=f"file:{rel}",
            title=path.stem,
            content=content,
            doc_type="file",
            metadata={"path": str(rel), "extension": path.suffix.lower(), "filename": path.name},
        )